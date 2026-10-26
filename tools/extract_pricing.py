"""Rebuild catalog data from FIXX_pricing_catalog.xlsx.

Run from the repo root after a price check:

    python3 tools/extract_pricing.py

The workbook must already have calculated values (open and save it in Excel
or LibreOffice first). This writes:

    src/seed/pricing/catalog.json
    tests/fixtures/quote_builder.json
"""
import datetime
import hashlib
import json
import sys

import openpyxl

src = sys.argv[1] if len(sys.argv) > 1 else "FIXX_pricing_catalog.xlsx"
wv = openpyxl.load_workbook(src, data_only=True)

TIER_ORDER = [
    "Budget", "Standard", "Premium",
    "1 t", "1.5 t", "2 t", "3 t",
    "≤60 m²", "61–100 m²", "101–150 m²",
]


def header_row(ws):
    hdr = []
    for cell in ws[1]:
        if cell.value is None:
            break
        hdr.append(cell.value)
    return hdr


def rows(name):
    ws = wv[name]
    hdr = header_row(ws)
    for r in ws.iter_rows(min_row=2, max_col=len(hdr)):
        if r[0].value not in (None, ""):
            yield dict(zip(hdr, [c.value for c in r]))


def col(r, prefix):
    for key, value in r.items():
        if key and str(key).startswith(prefix):
            return value
    raise KeyError(prefix)


def iso(d):
    return d.date().isoformat() if isinstance(d, (datetime.datetime, datetime.date)) else None


def num(v):
    return v if isinstance(v, (int, float)) else None


def tier_key(name):
    try:
        return TIER_ORDER.index(name)
    except ValueError:
        return len(TIER_ORDER)


items = {}
for r in rows("Prices"):
    name = r["Item (English)"] or ""
    notes = str(r.get("Notes") or "")
    items[r["Item_ID"]] = dict(
        id=r["Item_ID"], category=r["Category"], name=name, spanish=r["Spanish_Name"], spec=r["Spec"],
        unit=r["Unit / Pack"], store=r["Store"], price=num(r["Current_Price_MXN"]), iva=r["IVA_Status"],
        checkDate=iso(r["Check_Date"]), flag=r["Flag"] or "none",
        homeownerChoice=("homeowner reference" in name.lower()) or ("Homeowner-choice" in notes))

tasks = {}
for r in rows("Tasks"):
    trigger = (r.get("Review_Trigger (base / as-needed / none)") or "none")
    if trigger not in ("base", "as-needed", "none"):
        trigger = "none"
    t = tasks.setdefault(r["Task_ID"], dict(
        id=r["Task_ID"], category=r["Category"], job=r["Job"],
        jobEs=r.get("Job_Name_ES") or r["Job"],
        jobKo=r.get("Job_Name_KO (machine-translated, needs native review)") or r.get("Job_Name_ES") or r["Job"],
        lines=[]))
    if r.get("Job_Name_ES"):
        t["jobEs"] = r["Job_Name_ES"]
    if r.get("Job_Name_KO (machine-translated, needs native review)"):
        t["jobKo"] = r["Job_Name_KO (machine-translated, needs native review)"]
    t["lines"].append(dict(
        id=r["Line_ID"], tier=r["Tier"], itemId=r["Item_ID"], qty=r["Qty_Used"], basis=r["Qty_Basis"],
        unitPrice=num(r["Unit_Price_MXN"]), type=r["Line_Type"], flag=r["Flag (lookup)"] or "none",
        reviewTrigger=trigger))
for t in tasks.values():
    t["tiers"] = sorted({line["tier"] for line in t["lines"]}, key=tier_key)

mw = wv["Min_Wage_Floor"]
daily = {}
floors = {}
mode = "daily"
for r in range(2, mw.max_row + 1):
    a = mw.cell(r, 1).value
    if a == "Labor_Key":
        mode = "floor"
        continue
    if not isinstance(a, str):
        continue
    if mode == "daily" and isinstance(mw.cell(r, 2).value, (int, float)):
        daily[a] = mw.cell(r, 2).value
    if mode == "floor" and "|" in a:
        floors[a] = dict(
            trade=mw.cell(r, 5).value, dailyMin=mw.cell(r, 6).value, estJobDays=mw.cell(r, 7).value,
            billableDays=mw.cell(r, 8).value, floor=mw.cell(r, 9).value)

labor = {}
for r in rows("Labor_Rates"):
    f = floors.get(r["Labor_Key"], {})
    labor[r["Labor_Key"]] = dict(
        key=r["Labor_Key"], taskId=r["Task_ID"], tierScope=r["Tier_Scope"], trade=r["Trade"], basis=r["Labor_Basis"],
        rateLow=num(r["Rate_Low"]), rateTypical=num(r["Rate_Typical"]), rateHigh=num(r["Rate_High"]), refQty=num(r["Ref_Qty"]),
        typicalRef=num(r["Typical_Labor_Ref_Job"]), confidence=r["Confidence"], flag=r["Labor_Flag"] or "none",
        checkedDate=iso(r["Checked_Date"]), estJobDays=f.get("estJobDays"), dailyMin=f.get("dailyMin"), floor=f.get("floor"))

urgency = [dict(level=r["Urgency_Level"], multiplier=r["Multiplier (input)"], uiLevel=r["UI_Level"], uiOption=r["UI_Option"])
           for r in rows("Urgency") if r["Urgency_Level"] in ("STANDARD", "MEDIUM", "HIGH", "EMERGENCY")]
diagnosis = [dict(trade=r["Trade"], low=num(r["Fee_Low"]), typical=num(r["Fee_Typical"]), high=num(r["Fee_High"]), flag=r["Flag"])
             for r in rows("Diagnosis_Fees") if r["Trade"] and not str(r["Trade"]).startswith("FIXX POLICY")]
idmap = {r["Old_Task_ID"]: r["Final_Task_ID"] for r in rows("Task_ID_Map") if r.get("Old_Task_ID")}

dates = [i["checkDate"] for i in items.values() if i["checkDate"]] + [row["checkedDate"] for row in labor.values() if row["checkedDate"]]
body = dict(tasks=tasks, items=items, labor=labor, dailyMin=daily, urgency=urgency, diagnosis=diagnosis, taskIdMap=idmap)
digest = hashlib.sha1(json.dumps(body, sort_keys=True, default=str).encode()).hexdigest()[:8]
catalog = dict(version=f"{max(dates)}-{digest}", **body)
json.dump(catalog, open("src/seed/pricing/catalog.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
open("src/seed/pricing/catalog.json", "a", encoding="utf-8").write("\n")

qb = []
for r in rows("Quote_Builder"):
    qb.append(dict(
        taskId=r["Task_ID"], tier=r["Tier"], urgency=r["Urgency"],
        materials=num(col(r, "Materials_MXN")), laborKey=r["Labor_Key"],
        typicalLabor=num(col(r, "Typical_Labor_MXN")), multiplier=r["Urgency_Multiplier"],
        labor=num(col(r, "Labor_MXN")), total=num(col(r, "Customer_Total_MXN")),
        payout=num(col(r, "Pro_Payout_MXN")), status=r["Status"],
        reviewStatus=r.get("Review_Status"),
        jobEs=r.get("Job_Name_ES"),
        jobKo=col(r, "Job_Name_KO")))
json.dump(qb, open("tests/fixtures/quote_builder.json", "w", encoding="utf-8"), ensure_ascii=False)
open("tests/fixtures/quote_builder.json", "a", encoding="utf-8").write("\n")
print(f"catalog {catalog['version']}: {len(tasks)} tasks, {len(items)} items, {len(labor)} labor keys, {len(daily)} wage trades; fixture rows: {len(qb)}")
